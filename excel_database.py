from openpyxl import Workbook, load_workbook
import os


def save_card(card, card_number, image_path):

    filename = "pokemon_collection.xlsx"

    # Check if our Excel file already exists
    if os.path.exists(filename):
        workbook = load_workbook(filename)
        sheet = workbook.active

    else:
        # Create a new Excel file
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pokemon Collection"

        # Create column headers
        sheet.append([
            "Pokemon",
            "Card Number",
            "Set",
            "Rarity",
            "HP",
            "Illustrator",
            "Database ID",
            "Image Path"
        ])

    # Add the scanned card as a new row
    sheet.append([
        card["name"],
        card_number,
        card["set"]["name"],
        card.get("rarity", "Unknown"),
        card.get("hp", "Unknown"),
        card.get("illustrator", "Unknown"),
        card["id"],
        image_path
    ])

    # Save the Excel file
    workbook.save(filename)

    print("Card saved to Excel!")